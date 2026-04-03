#!/usr/bin/env python3
"""
csv2xlsx - CSV and XLSX file converter

Usage:
    python csv2xlsx.py input.csv -o output.xlsx
    python csv2xlsx.py input.xlsx -o output.csv

Args:
    input: Input file path (csv or xlsx)
    -o, --output: Output file path
    -e, --encoding: Input file encoding (auto-detected if not specified)
"""

import sys
import argparse
from pathlib import Path


def detect_encoding(file_path: Path) -> str:
    import chardet

    with open(file_path, "rb") as f:
        raw_data = f.read()
    result = chardet.detect(raw_data)
    encoding = result["encoding"]
    confidence = result["confidence"]
    print(f"Detected encoding: {encoding} (confidence: {confidence:.2%})")
    if encoding is None:
        raise ValueError("Could not detect file encoding")
    return encoding


def csv_to_xlsx(input_path: Path, output_path: Path, encoding: str = None):
    import csv
    from openpyxl import Workbook
    from openpyxl.styles import Alignment

    if encoding is None:
        encoding = detect_encoding(input_path)

    wb = Workbook()
    ws = wb.active

    with open(input_path, "r", encoding=encoding, newline="") as f:
        reader = csv.reader(f)
        for row_idx, row in enumerate(reader, 1):
            for col_idx, value in enumerate(row, 1):
                cell = ws.cell(row=row_idx, column=col_idx, value=value)
                cell.alignment = Alignment(wrap_text=True)

    wb.save(output_path)
    print(f"Converted: {input_path} -> {output_path}")


def xlsx_to_csv(input_path: Path, output_path: Path):
    import csv
    from openpyxl import load_workbook

    wb = load_workbook(input_path, read_only=True)
    ws = wb.active

    with open(output_path, "w", encoding="utf-8-sig", newline="") as f:
        writer = csv.writer(f)
        for row in ws.iter_rows(values_only=True):
            writer.writerow(row)

    wb.close()
    print(f"Converted: {input_path} -> {output_path}")


def main():
    parser = argparse.ArgumentParser(description="Convert between CSV and XLSX files")
    parser.add_argument("input", type=Path, help="Input file (csv or xlsx)")
    parser.add_argument("-o", "--output", type=Path, required=True, help="Output file path")
    parser.add_argument("-e", "--encoding", help="Input file encoding (auto-detected if not specified)")
    args = parser.parse_args()

    input_ext = args.input.suffix.lower()
    output_ext = args.output.suffix.lower()

    if not args.input.exists():
        print(f"Error: Input file not found: {args.input}")
        sys.exit(1)

    encoding = args.encoding

    if input_ext == ".csv" and output_ext == ".xlsx":
        csv_to_xlsx(args.input, args.output, encoding)
    elif input_ext == ".xlsx" and output_ext == ".csv":
        xlsx_to_csv(args.input, args.output)
    else:
        print(f"Error: Unsupported conversion from {input_ext} to {output_ext}")
        print("Supported conversions: CSV <-> XLSX")
        sys.exit(1)


if __name__ == "__main__":
    main()
